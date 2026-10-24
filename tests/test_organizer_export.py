from io import BytesIO
import copy
from openpyxl import Workbook, load_workbook
from fastapi.testclient import TestClient
import pytest
from backend import app as module, auth
from backend.db import Database
from backend.sample import sample_state
from backend.organizer_export import organizer_source, organizer_workbook, ORG_HEADERS, POSITION_HEADERS


def source_workbook():
    w=Workbook();s=w.active;s.title='Org Structure'
    s.append(['Org Chart']);s.append([]);s.append([]);s.append(ORG_HEADERS)
    s.append(['0000','','Corporation','Root owner','Purpose',0])
    s.append(['1000','0000','Technology','Owner','Original purpose',1])
    s=w.create_sheet('Position Control');s.append(['Position Control']);s.append([]);s.append(POSITION_HEADERS)
    for data in [
        {'Unique Company ID':'030','Position ID':'0042','Name':'Original name','Company':'DTI','Department ID':'1000','Position ID Status':'Active','Recruiter':'Recruiter value','Priority':'High','Tier':'Dependent','PID Approval Date':'2026-01-03'},
        {'Unique Company ID':'Unknown','Position ID':'0043','Company':'CZV','Department ID':'1000','Position ID Status':'GATED','Recruiter':'Source only recruiter'},
    ]:s.append([data.get(h,'') for h in POSITION_HEADERS])
    b=BytesIO();w.save(b);return b.getvalue()


@pytest.fixture
def client(monkeypatch,tmp_path):
    db=Database('sqlite:///'+str(tmp_path/'export.db'));db.save(sample_state(),0,'test','seed')
    monkeypatch.setattr(module,'DB',db);monkeypatch.setattr(auth,'MODE','development')
    with TestClient(module.app) as c: yield c


def read_rows(sheet,header):
    names=[c.value for c in sheet[header]]
    return [dict(zip(names,row)) for row in sheet.iter_rows(min_row=header+1,values_only=True)]


def import_source(client):
    r=client.post('/api/imports/preview',files={'file':('Org Chart Organizer.xlsx',source_workbook())})
    assert r.status_code==200,r.text
    preview=r.json();assert not preview['errors']
    assert 'organizerSource' not in preview
    r=client.post('/api/imports/'+preview['id']+'/commit',json={'acknowledgeWarnings':True,'confirmConflicts':True})
    assert r.status_code==200,r.text


def test_source_retained_on_commit_and_export_updates_current_values(client):
    import_source(client)
    saved=copy.deepcopy(module.DB.get_setting('organizerSource'))
    assert saved['positions'][0]['values']['Recruiter']=='Recruiter value'
    state,v=module.DB.read()
    unit=dict(state['OrganizationalUnits']['1000'],DepartmentPurpose='Updated purpose',DepartmentOwner='Updated owner')
    r=client.put('/api/planning/OrganizationalUnits/1000',json={'version':v,'record':unit});assert r.status_code==200,r.text
    state,v=module.DB.read();state['Employees']['030'].update(Name='Current name',ManagerID='001',SourceCompany='czinger',OrgUnitID='1000')
    module.DB.save(state,v,'test','fixture')
    before=module.DB.read()
    r=client.get('/api/workbooks/organizer');assert r.status_code==200,r.text
    assert 'Org-Chart-Organizer.xlsx' in r.headers['content-disposition']
    w=load_workbook(BytesIO(r.content),data_only=False)
    assert [c.value for c in w['Org Structure'][4]]==ORG_HEADERS
    assert [c.value for c in w['Position Control'][3]]==POSITION_HEADERS
    org={r['Department ID']:r for r in read_rows(w['Org Structure'],4)}
    assert org['1000']['Department Purpose']=='Updated purpose' and org['1000']['Department Owner']=='Updated owner'
    rows=read_rows(w['Position Control'],3)
    p=next(r for r in rows if r['Unique Company ID']=='030')
    assert p['Name']=='Current name' and p['Company']=='CZV' and p['Department ID']=='1000'
    assert p['Position ID']=='0042' and p['Recruiter']=='Recruiter value' and p['Priority']=='High'
    assert p['Tier']=='Dependent'
    assert p['Direct Manager']==state['Employees']['001']['Name']
    assert p['PID Approval Date'].date().isoformat()=='2026-01-03'
    unknown=next(r for r in rows if r['Unique Company ID']=='Unknown')
    assert unknown['Position ID Status']=='GATED' and unknown['Recruiter']=='Source only recruiter'
    assert 'tblOrgStructure' in w['Org Structure'].tables and 'tblPositionControl' in w['Position Control'].tables
    assert 'tblPeople' in w['People'].tables and 'tblOrgUnits' in w['Org Units'].tables
    current=read_rows(w['People'],1)
    assert len(current)==len(state['Employees']) and len({p['EmployeeID'] for p in current})==len(current)
    assert all(p['OrgUnitID'] in org for p in current)
    assert not any('Compensation' in sheet for sheet in w.sheetnames)
    assert module.DB.read()==before and module.DB.get_setting('organizerSource')==saved


def test_formula_like_values_are_literal_and_missing_fields_are_reported():
    state=sample_state();state['Employees']['030']['Name']='=HYPERLINK("https://example.test")'
    state['OrganizationalUnits']['company']['DepartmentPurpose']='=1+1'
    w=load_workbook(BytesIO(organizer_workbook(state,organizer_source(source_workbook()))),data_only=False)
    for s in w:
        assert all(c.data_type!='f' for row in s for c in row)
    assert any(c.value=='=1+1' and c.data_type=='s' for row in w['Org Structure'] for c in row)
    notes=read_rows(w['Export Notes'],1)
    assert any('not supplied' in r['Review note'] for r in notes)
    assert any(r['Match status']=='Source only' for r in read_rows(w['Position Links'],1))


def test_vacancies_keep_source_approval_states_without_inventing_new_ones():
    state=sample_state();base=copy.deepcopy(state['Employees']['030'])
    state['Employees']['77-Open']=dict(base,EmployeeID='77-Open',Name='Open Role',EmploymentStatus='vacant',FTE='0')
    state['Employees']['new-open']=dict(base,EmployeeID='new-open',Name='Open Role',EmploymentStatus='vacant',FTE='0')
    source={'headers':POSITION_HEADERS,'positions':[{'sourceRow':4,'values':{'Unique Company ID':'DTI','Position ID':'77','Position ID Status':'HOLD - Approved'}}]}
    w=load_workbook(BytesIO(organizer_workbook(state,source)))
    rows=read_rows(w['Position Control'],3)
    assert rows[0]['Position ID Status']=='HOLD - Approved' and rows[0]['Name'] is None
    new=next(r for r in rows if r['Unique Company ID']=='new-open')
    assert new['Position ID Status'] is None
    links=read_rows(w['Position Links'],1)
    assert len([r for r in links if r['EmployeeID']=='77-Open'])==1


@pytest.mark.parametrize('role',['viewer','owner'])
def test_export_permission_is_enforced_on_server(client,role):
    client.post('/api/dev-session',json={'role':role})
    assert client.get('/api/workbooks/organizer').status_code==403


def test_stale_import_does_not_replace_source_snapshot(client):
    import_source(client);saved=copy.deepcopy(module.DB.get_setting('organizerSource'))
    r=client.post('/api/imports/preview',files={'file':('new.xlsx',source_workbook())}).json()
    state,v=module.DB.read();module.DB.save(state,v,'test','intervening edit')
    response=client.post('/api/imports/'+r['id']+'/commit',json={'acknowledgeWarnings':True,'confirmConflicts':True})
    assert response.status_code==409
    assert module.DB.get_setting('organizerSource')==saved


def test_nonstandard_sensitive_columns_are_not_exported():
    w=load_workbook(BytesIO(source_workbook()));s=w['Position Control']
    s.cell(3,33,'Annual Salary');s.cell(4,33,'secret compensation value')
    s.cell(3,34,'Performance Score');s.cell(4,34,'secret performance value')
    b=BytesIO();w.save(b)
    source=organizer_source(b.getvalue())
    assert source['omittedColumnCount']==2
    exported=load_workbook(BytesIO(organizer_workbook(sample_state(),source)))
    values=[str(c.value) for sheet in exported for row in sheet for c in row]
    assert 'Annual Salary' not in values and 'Performance Score' not in values
    assert not any('secret compensation' in v or 'secret performance' in v for v in values)


def test_invalid_source_date_is_preserved_and_decimal_tiers_export_as_numbers():
    state=sample_state();state['OrganizationalUnits']['company']['Tier']='0.0'
    source=organizer_source(source_workbook())
    source['positions'][0]['values']['Need Date']='2026-02-30'
    w=load_workbook(BytesIO(organizer_workbook(state,source)))
    root=next(r for r in read_rows(w['Org Structure'],4) if r['Department ID']=='company')
    assert root['Tier']==0
    assert read_rows(w['Position Control'],3)[0]['Need Date']=='2026-02-30'
