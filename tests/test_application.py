import copy, io, json, time
from datetime import date
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from backend import app as module
from backend import auth
from backend.db import Database, StaleWrite
from backend.domain import validate, payroll, budget_metrics, kr_progress, descendants
from backend.sample import sample_state, workbook_bytes


@pytest.fixture
def client(monkeypatch, tmp_path):
    db = Database("sqlite:///" + str(tmp_path / "test.db"))
    db.save(sample_state(), 0, "test", "seed", setting=("synthetic", True))
    monkeypatch.setattr(module, "DB", db)
    monkeypatch.setattr(auth, "MODE", "development")
    with TestClient(module.app) as c:
        yield c


def role(c, r):
    assert c.post("/api/dev-session", json={"role": r}).status_code == 200


def preview(c, state=None, content=None, mode="incremental", mapping=None):
    return c.post(
        "/api/imports/preview",
        files={
            "file": (
                "test.xlsx",
                content or workbook_bytes(state or sample_state()),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        data={"mode": mode, "mapping": json.dumps(mapping or {})},
    )


def commit(c, p, **kwargs):
    return c.post("/api/imports/" + p["id"] + "/commit", json=kwargs)


def edit(c, sheet, record):
    key = module.COLUMNS[sheet][0]
    v = c.get("/api/organization").json()["version"]
    return c.put(
        f"/api/planning/{sheet}/{record[key]}", json={"version": v, "record": record}
    )


def test_sample_contract_and_canonical_totals(client):
    state = sample_state()
    assert validate(state) == ([], [])
    assert len(state["Employees"]) == 30
    totals = payroll(state, set(state["Employees"]), date.today().isoformat(), "USD")
    expected = sum(
        int(c["AnnualBasePay"])
        + int(c["AnnualTargetBonus"])
        + int(c["AnnualEmployerCosts"])
        for c in state["Compensation"].values()
    )
    assert totals == {
        "headcount": 30,
        "fte": 29.8,
        "payroll": expected,
        "missingPay": 0,
        "otherCurrency": 0,
        "complete": True,
    }
    r = client.get("/api/organization").json()
    ceo = next(e for e in r["employees"] if e["EmployeeID"] == "001")
    assert ceo["directReports"] == 3 and ceo["downstreamReports"] == 29
    assert ceo["metrics"]["payroll"] == expected
    assert r["summary"]["payroll"] == expected


def test_budget_envelopes_not_double_counted():
    state = sample_state()
    r = budget_metrics(
        state, "company", date.today().year, "USD", date.today().isoformat()
    )
    assert r["envelope"] == 7500000 and r["childEnvelopes"] == 6900000
    assert r["unallocated"] == 300000 and r["directAllocation"] == 300000
    assert r["direct"]["payroll"] == 345000 and r["directHeadroom"] == -45000
    assert r["total"]["payroll"] != sum(
        budget_metrics(state, u, date.today().year, "USD", date.today().isoformat())[
            "total"
        ]["payroll"]
        for u in state["OrganizationalUnits"]
    )


def test_fte_not_multiplied_twice_and_missing_pay_not_zero():
    s = sample_state()
    r = payroll(s, {"021"}, date.today().isoformat(), "USD")
    assert r["fte"] == 0.8 and r["payroll"] == 140000
    del s["Compensation"]["021@2025-01-01"]
    r = payroll(s, {"021"}, date.today().isoformat(), "USD")
    assert r["missingPay"] == 1 and r["complete"] is False
    s["Compensation"]["021@2025-01-01"] = dict(
        sample_state()["Compensation"]["021@2025-01-01"], Currency="EUR"
    )
    r = payroll(s, {"021"}, date.today().isoformat(), "USD")
    assert r["otherCurrency"] == 1 and not r["complete"]


def test_distinct_hierarchies():
    s = sample_state()
    s["Employees"]["030"]["ManagerID"] = "003"
    reporting = descendants(s["Employees"], "ManagerID", "003")
    ops = descendants(s["OrganizationalUnits"], "ParentOrgUnitID", "ops")
    assert "030" in reporting and s["Employees"]["030"]["OrgUnitID"] not in ops


def test_import_idempotent_and_leading_zero_ids(client):
    p = preview(client).json()
    assert not p["errors"]
    assert p["changes"] == []
    assert commit(client, p).status_code == 200
    p = preview(client).json()
    assert commit(client, p).status_code == 200
    data = client.get("/api/organization").json()
    assert len(data["employees"]) == 30
    assert data["employees"][0]["EmployeeID"] == "001"


def test_move_stable_id_no_duplicate_and_audit(client):
    s = sample_state()
    s["Employees"]["030"]["ManagerID"] = "003"
    s["Employees"]["030"]["OrgUnitID"] = "people"
    p = preview(client, s).json()
    change = next(x for x in p["changes"] if x["id"] == "030")
    assert change["reportingChange"] and change["unitMove"]
    assert commit(client, p).status_code == 200
    data = client.get("/api/organization").json()
    assert len(data["employees"]) == 30
    assert (
        next(e for e in data["employees"] if e["EmployeeID"] == "030")["ManagerID"]
        == "003"
    )
    events = client.get("/api/audit").json()
    assert events[0]["details"]["changes"][0]["reportingChange"]


@pytest.mark.parametrize(
    "mutation,expected",
    [
        (lambda s: s["Employees"]["030"].update(ManagerID="missing"), "ManagerID"),
        (lambda s: s["Employees"]["001"].update(ManagerID="030"), "ManagerID"),
        (lambda s: s["Employees"]["030"].update(FTE="2"), "FTE"),
        (lambda s: s["Employees"]["030"].update(FTE="NaN"), "FTE"),
        (lambda s: s["Employees"]["030"].update(StartDate="tomorrow"), "StartDate"),
        (
            lambda s: s["Compensation"]["030@2025-01-01"].update(AnnualBasePay="-1"),
            "AnnualBasePay",
        ),
        (
            lambda s: s["Compensation"]["030@2025-01-01"].update(Currency="banana"),
            "Currency",
        ),
        (
            lambda s: s["OrganizationalUnits"]["eng"].update(
                ParentOrgUnitID="software"
            ),
            "ParentOrgUnitID",
        ),
    ],
)
def test_invalid_import_keeps_previous_dataset(client, mutation, expected):
    before = client.get("/api/organization").json()
    s = sample_state()
    mutation(s)
    p = preview(client, s).json()
    assert any(e["field"] == expected for e in p["errors"])
    assert commit(client, p).status_code == 422
    assert client.get("/api/organization").json() == before


def test_duplicates_formulas_and_column_mapping(client):
    wb = load_workbook(io.BytesIO(workbook_bytes(sample_state())))
    ws = wb["Employees"]
    ws.append([c.value for c in ws[2]])
    buf = io.BytesIO()
    wb.save(buf)
    p = preview(client, content=buf.getvalue()).json()
    assert any("Duplicate ID" in e["message"] for e in p["errors"])
    ws.delete_rows(ws.max_row)
    ws["B1"] = "Full Name"
    ws["B2"] = "=1+1"
    buf = io.BytesIO()
    wb.save(buf)
    p = preview(
        client, content=buf.getvalue(), mapping={"Employees": {"Name": "Full Name"}}
    ).json()
    assert any("literal" in e["message"] for e in p["errors"])
    ws["B2"] = "Lukas Czinger"
    buf = io.BytesIO()
    wb.save(buf)
    p = preview(
        client, content=buf.getvalue(), mapping={"Employees": {"Name": "Full Name"}}
    ).json()
    assert not p["errors"]


def test_overlap_and_duplicate_budget_validation(client):
    s = sample_state()
    s["Compensation"]["extra"] = dict(
        s["Compensation"]["001@2025-01-01"], EffectiveFrom="2026-01-01"
    )
    p = preview(client, s).json()
    assert any("Overlapping" in e["message"] for e in p["errors"])
    s = sample_state()
    r = next(iter(s["Budgets"].values()))
    s["Budgets"]["new"] = dict(r, BudgetID="new")
    p = preview(client, s).json()
    assert any("Duplicate budget" in e["message"] for e in p["errors"])


def test_incremental_missing_rows_and_snapshot_confirmation(client):
    s = sample_state()
    del s["Employees"]["030"]
    p = preview(client, s).json()
    assert not any(c["kind"] == "deactivation" for c in p["changes"])
    assert commit(client, p).status_code == 200
    p = preview(client, s, mode="snapshot").json()
    assert len([c for c in p["changes"] if c["kind"] == "deactivation"]) == 1
    assert commit(client, p).status_code == 409
    assert commit(client, p, confirmDeactivations=True).status_code == 200
    assert len(client.get("/api/organization").json()["employees"]) == 29
    state, _ = module.DB.read()
    assert state["Employees"]["030"]["EmploymentStatus"] == "inactive"
    assert "030@2025-01-01" in state["Compensation"]


def test_multiple_roots_require_review(client):
    s = sample_state()
    s["Employees"]["030"]["ManagerID"] = ""
    p = preview(client, s).json()
    assert p["warnings"]
    assert commit(client, p).status_code == 409
    assert commit(client, p, acknowledgeWarnings=True).status_code == 200


def test_optional_sheets_and_template(client):
    wb = load_workbook(io.BytesIO(workbook_bytes(sample_state())))
    for sheet in ["Compensation", "Budgets", "Objectives", "KeyResults"]:
        del wb[sheet]
    buf = io.BytesIO()
    wb.save(buf)
    p = preview(client, content=buf.getvalue()).json()
    assert not p["errors"]
    assert commit(client, p).status_code == 200
    template = client.get("/api/workbooks/template")
    assert template.status_code == 200
    wb = load_workbook(io.BytesIO(template.content))
    assert wb["Employees"].max_row == 1


def test_app_owned_plans_survive_conflicting_imports(client):
    obj = sample_state()["Objectives"]["o-eng"].copy()
    obj["ObjectiveText"] = "App-owned strategy"
    assert edit(client, "Objectives", obj).status_code == 200
    kr = sample_state()["KeyResults"]["kr1"].copy()
    kr["CurrentValue"] = "9"
    assert edit(client, "KeyResults", kr).status_code == 200
    b = sample_state()["Budgets"][f"budget-eng-{date.today().year}"].copy()
    b["DirectAllocation"] = "290000"
    assert edit(client, "Budgets", b).status_code == 200
    unit = {
        "OrgUnitID": "new-team",
        "OrgUnitName": "New team",
        "ParentOrgUnitID": "eng",
        "UnitType": "subdepartment",
        "OwnerEmployeeID": "002",
    }
    assert edit(client, "OrganizationalUnits", unit).status_code == 200
    p = preview(client).json()
    assert len(p["conflicts"]) == 3
    assert commit(client, p).status_code == 409
    assert commit(client, p, confirmConflicts=True).status_code == 200
    s, _ = module.DB.read()
    assert s["Objectives"]["o-eng"]["ObjectiveText"] == "App-owned strategy"
    assert s["KeyResults"]["kr1"]["CurrentValue"] == "9"
    assert "new-team" in s["OrganizationalUnits"]
    assert s["Budgets"][b["BudgetID"]]["DirectAllocation"] == "290000"


def test_restore_preserves_later_app_work(client):
    p = preview(client).json()
    assert commit(client, p).status_code == 200
    obj = sample_state()["Objectives"]["o-eng"].copy()
    obj["ObjectiveText"] = "Later planning edit"
    assert edit(client, "Objectives", obj).status_code == 200
    r = client.post(f"/api/imports/{p['id']}/restore-preview").json()
    assert r["restore"]
    assert commit(client, r, confirmConflicts=True).status_code == 200
    assert (
        module.DB.read()[0]["Objectives"]["o-eng"]["ObjectiveText"]
        == "Later planning edit"
    )


def test_stale_preview_and_edits_rejected(client):
    p = preview(client).json()
    obj = sample_state()["Objectives"]["o-eng"]
    assert edit(client, "Objectives", obj).status_code == 200
    assert commit(client, p).status_code == 409
    assert (
        client.put(
            "/api/planning/Objectives/o-eng", json={"version": 1, "record": obj}
        ).status_code
        == 409
    )


@pytest.mark.parametrize(
    "path",
    [
        "/api/people/001/compensation",
        "/api/workbooks/export",
        "/api/audit",
        "/api/imports",
        "/api/integration",
    ],
)
def test_viewer_cannot_read_restricted_endpoints(client, path):
    role(client, "viewer")
    assert client.get(path).status_code == 403


def test_viewer_cannot_mutate_or_receive_pay(client):
    role(client, "viewer")
    r = client.get("/api/organization").json()
    assert r["summary"]["payroll"] is None
    assert all(e["metrics"]["payroll"] is None for e in r["employees"])
    assert all(
        u["metrics"]["total"]["payroll"] is None and "budget" not in u["metrics"]
        for u in r["units"]
    )
    assert preview(client).status_code == 403
    assert (
        edit(client, "Objectives", sample_state()["Objectives"]["o-eng"]).status_code
        == 403
    )


def test_owner_scope_and_subtraction_protection(client):
    role(client, "owner")
    r = client.get("/api/organization").json()
    assert r["summary"]["payroll"] is not None
    allowed = descendants(
        sample_state()["OrganizationalUnits"], "ParentOrgUnitID", "eng"
    )
    assert all(e["OrgUnitID"] in allowed for e in r["employees"])
    assert all(e["metrics"]["payroll"] is None for e in r["employees"])
    assert all(u["metrics"]["direct"]["payroll"] is None for u in r["units"])
    assert all(
        u["metrics"]["total"]["payroll"] is None
        for u in r["units"]
        if u["OrgUnitID"] != "eng"
    )
    assert client.get("/api/people/002/compensation").status_code == 403
    assert client.get("/api/workbooks/export").status_code == 403
    assert client.get("/api/organization?as_of=2025-01-01").status_code == 403
    assert (
        edit(client, "Objectives", sample_state()["Objectives"]["o-ops"]).status_code
        == 403
    )
    foreign = sample_state()["KeyResults"]["kr2"]
    foreign["ObjectiveID"] = "o-eng"
    assert edit(client, "KeyResults", foreign).status_code == 403


def test_owner_small_scope_suppressed(client, monkeypatch):
    monkeypatch.setitem(
        auth.DEV_ROLES, "owner", auth.Principal("demo-owner", "owner", "brand")
    )
    role(client, "owner")
    r = client.get("/api/organization").json()
    assert r["summary"]["headcount"] < 5
    assert r["summary"]["payroll"] is None


def test_owner_budget_permissions_and_new_unit(client):
    role(client, "owner")
    s = sample_state()
    b = s["Budgets"][f"budget-eng-{date.today().year}"].copy()
    b["ApprovedEnvelope"] = "3000000"
    assert edit(client, "Budgets", b).status_code == 403
    b = s["Budgets"][f"budget-vehicle-{date.today().year}"].copy()
    b["ApprovedEnvelope"] = "2000000"
    assert edit(client, "Budgets", b).status_code == 422
    unit = {
        "OrgUnitID": "new-team",
        "OrgUnitName": "New team",
        "ParentOrgUnitID": "eng",
        "UnitType": "subdepartment",
        "OwnerEmployeeID": "002",
    }
    assert edit(client, "OrganizationalUnits", unit).status_code == 200
    unit = dict(unit, OrgUnitID="bad-team", ParentOrgUnitID="ops")
    assert edit(client, "OrganizationalUnits", unit).status_code == 403


def test_okr_increase_decrease_equal_missing_stale():
    s = sample_state()
    assert (
        kr_progress(s["KeyResults"]["kr1"], date.today().isoformat())["progress"] == 70
    )
    assert (
        kr_progress(s["KeyResults"]["kr2"], date.today().isoformat())["progress"]
        == 53.3
    )
    row = dict(s["KeyResults"]["kr1"], TargetValue="0", LastUpdated="2025-01-01")
    r = kr_progress(row, date.today().isoformat())
    assert r["progress"] is None and r["stale"]
    row["CurrentValue"] = ""
    assert kr_progress(row, date.today().isoformat())["reason"] == "Missing measurement"


def test_persistence_after_database_reopen(tmp_path):
    url = "sqlite:///" + str(tmp_path / "persistent.db")
    db = Database(url)
    db.save(sample_state(), 0, "test", "seed")
    db.engine.dispose()
    again = Database(url)
    state, v = again.read()
    assert v == 1 and len(state["Employees"]) == 30


def test_unconfigured_integration_preserves_dataset(client):
    before = client.get("/api/organization").json()
    assert not client.get("/api/integration").json()["configured"]
    assert client.post("/api/integration/refresh").status_code == 409
    assert client.get("/api/organization").json() == before


def test_production_rejects_development_auth(monkeypatch):
    monkeypatch.setenv("APP_ENV", "production")
    monkeypatch.setattr(auth, "MODE", "development")
    with pytest.raises(RuntimeError, match="Production requires"):
        auth.validate_config()


def test_cross_origin_mutation_rejected(client):
    assert (
        client.post(
            "/api/dev-session",
            json={"role": "admin"},
            headers={"Origin": "https://evil.example"},
        ).status_code
        == 403
    )


def test_1000_people_functional_performance(client):
    s = sample_state(1000)
    assert validate(s) == ([], [])
    _, v = module.DB.read()
    module.DB.save(s, v, "test", "large-sample")
    start = time.perf_counter()
    r = client.get("/api/organization")
    elapsed = time.perf_counter() - start
    assert r.status_code == 200
    data = r.json()
    assert len(data["employees"]) == 1000 and data["summary"]["headcount"] == 1000
    ceo = next(e for e in data["employees"] if e["EmployeeID"] == "001")
    assert ceo["downstreamReports"] == 999
    print(f"\n1000-person API: {elapsed:.3f}s, {len(r.content)} bytes")
    assert elapsed < 10


def test_historical_headcount_after_snapshot_deactivation(client):
    from datetime import timedelta

    s = sample_state()
    del s["Employees"]["030"]
    p = preview(client, s, mode="snapshot").json()
    assert commit(client, p, confirmDeactivations=True).status_code == 200
    before = (date.today() - timedelta(days=1)).isoformat()
    data = client.get("/api/organization?as_of=" + before).json()
    assert len(data["employees"]) == 30
    assert len(client.get("/api/organization").json()["employees"]) == 29


def test_owner_cannot_skip_unbudgeted_ancestor(client):
    role(client, "owner")
    for ident, parent in [("middle", "eng"), ("leaf", "middle")]:
        assert (
            edit(
                client,
                "OrganizationalUnits",
                {
                    "OrgUnitID": ident,
                    "OrgUnitName": ident,
                    "ParentOrgUnitID": parent,
                    "UnitType": "subdepartment",
                    "OwnerEmployeeID": "002",
                },
            ).status_code
            == 200
        )
    r = edit(
        client,
        "Budgets",
        {
            "BudgetID": "hidden-envelope",
            "OrgUnitID": "leaf",
            "BudgetYear": str(date.today().year),
            "Currency": "USD",
            "ApprovedEnvelope": "9000000",
            "DirectAllocation": "9000000",
        },
    )
    assert r.status_code == 422 and "parent unit" in r.text


def test_owner_company_alignment_is_read_only(client):
    role(client, "owner")
    r = client.get("/api/organization").json()
    company = next(o for o in r["objectives"] if o["ObjectiveID"] == "o-company")
    assert company["readOnly"] and not company["OwnerEmployeeID"]
    assert (
        edit(
            client, "Objectives", sample_state()["Objectives"]["o-company"]
        ).status_code
        == 403
    )
    own = sample_state()["Objectives"]["o-eng"]
    own["ObjectiveText"] = "Keep public company alignment"
    assert edit(client, "Objectives", own).status_code == 200


def test_planning_rejects_formula_injection(client):
    obj = sample_state()["Objectives"]["o-eng"]
    obj["ObjectiveText"] = '=HYPERLINK("https://example.invalid")'
    assert edit(client, "Objectives", obj).status_code == 422


def test_entra_signed_token_access_and_audience(client, monkeypatch):
    import jwt
    from cryptography.hazmat.primitives.asymmetric import rsa
    from types import SimpleNamespace

    key = rsa.generate_private_key(public_exponent=65537, key_size=2048)
    monkeypatch.setattr(auth, "MODE", "entra")
    for k, v in {
        "ENTRA_TENANT_ID": "tenant-test",
        "ENTRA_AUDIENCE": "api-test",
        "ENTRA_ACCESS_MAP": '{"person-test":{"role":"owner","unit":"eng"}}',
        "ENTRA_CLIENT_ID": "spa-test",
        "ENTRA_SCOPE": "api://api-test/access_as_user",
    }.items():
        monkeypatch.setenv(k, v)

    class Keys:
        def __init__(self, *args):
            pass

        def get_signing_key_from_jwt(self, *args):
            return SimpleNamespace(key=key.public_key())

    monkeypatch.setattr(jwt, "PyJWKClient", Keys)
    claims = {
        "scp": "access_as_user",
        "oid": "person-test",
        "iss": "https://login.microsoftonline.com/tenant-test/v2.0",
        "aud": "api-test",
        "iat": int(time.time()),
        "exp": int(time.time()) + 600,
    }
    token = jwt.encode(claims, key, algorithm="RS256")
    headers = {"Authorization": "Bearer " + token}
    assert client.get("/api/organization").status_code == 401
    assert client.get("/api/session", headers=headers).json()["role"] == "owner"
    assert (
        client.get("/api/people/001/compensation", headers=headers).status_code == 403
    )
    assert client.post("/api/dev-session", json={"role": "admin"}).status_code == 404
    bad = jwt.encode(dict(claims, aud="different-api"), key, algorithm="RS256")
    assert (
        client.get(
            "/api/session", headers={"Authorization": "Bearer " + bad}
        ).status_code
        == 401
    )
    unscoped = jwt.encode(dict(claims, scp="different-scope"), key, algorithm="RS256")
    assert (
        client.get(
            "/api/session", headers={"Authorization": "Bearer " + unscoped}
        ).status_code
        == 401
    )
    unknown = jwt.encode(dict(claims, oid="unassigned-user"), key, algorithm="RS256")
    assert (
        client.get(
            "/api/session", headers={"Authorization": "Bearer " + unknown}
        ).status_code
        == 401
    )


def test_graph_preview_commit_and_failure_preservation(client, monkeypatch):
    import httpx

    for k in [
        "GRAPH_TENANT_ID",
        "GRAPH_CLIENT_ID",
        "GRAPH_CLIENT_SECRET",
        "GRAPH_DRIVE_ID",
        "GRAPH_ITEM_ID",
    ]:
        monkeypatch.setenv(k, "test-only")
    content = workbook_bytes(sample_state())
    requests = []
    fail = [False]

    def handler(request):
        requests.append(request)
        if fail[0]:
            return httpx.Response(403, json={"error": "test-denial"})
        if "oauth2" in request.url.path:
            return httpx.Response(200, json={"access_token": "test-token"})
        if request.url.host == "graph.microsoft.com":
            return httpx.Response(
                200,
                json={
                    "eTag": "version-1",
                    "name": "corporate-test.xlsx",
                    "@microsoft.graph.downloadUrl": "https://tenant.sharepoint.test/file",
                },
            )
        assert "authorization" not in request.headers
        return httpx.Response(200, content=content)

    real = httpx.AsyncClient
    monkeypatch.setattr(
        httpx,
        "AsyncClient",
        lambda **kwargs: real(transport=httpx.MockTransport(handler), **kwargs),
    )
    before = client.get("/api/organization").json()
    r = client.post("/api/integration/refresh")
    assert r.status_code == 200, r.text
    status = r.json()
    assert status["state"] == "review-required"
    assert client.get("/api/organization").json() == before
    p = client.get("/api/imports/" + status["previewId"] + "/preview").json()
    assert commit(client, p).status_code == 200
    status = client.get("/api/integration").json()["status"]
    assert status["state"] == "applied" and status["lastSuccessfulRefresh"]
    assert client.post("/api/integration/refresh").json()["state"] == "applied"
    successful = client.get("/api/organization").json()
    fail[0] = True
    assert client.post("/api/integration/refresh").status_code == 502
    assert client.get("/api/organization").json() == successful
    assert client.get("/api/integration").json()["status"]["lastSuccessfulRefresh"]


def test_graph_stale_preview_is_regenerated(client, monkeypatch):
    import httpx

    for k in [
        "GRAPH_TENANT_ID",
        "GRAPH_CLIENT_ID",
        "GRAPH_CLIENT_SECRET",
        "GRAPH_DRIVE_ID",
        "GRAPH_ITEM_ID",
    ]:
        monkeypatch.setenv(k, "test-only")
    content = workbook_bytes(sample_state())

    def handler(request):
        if "oauth2" in request.url.path:
            return httpx.Response(200, json={"access_token": "test-token"})
        if request.url.host == "graph.microsoft.com":
            return httpx.Response(
                200,
                json={
                    "eTag": "version-1",
                    "name": "test.xlsx",
                    "@microsoft.graph.downloadUrl": "https://tenant.sharepoint.test/file",
                },
            )
        return httpx.Response(200, content=content)

    real = httpx.AsyncClient
    monkeypatch.setattr(
        httpx,
        "AsyncClient",
        lambda **kwargs: real(transport=httpx.MockTransport(handler), **kwargs),
    )
    first = client.post("/api/integration/refresh").json()["previewId"]
    obj = sample_state()["Objectives"]["o-eng"]
    obj["ObjectiveText"] = "Changed after preview"
    assert edit(client, "Objectives", obj).status_code == 200
    second = client.post("/api/integration/refresh").json()["previewId"]
    assert first != second


def test_explicit_tiers_and_duplicate_department_codes_roundtrip():
    state = sample_state()
    units = state["OrganizationalUnits"]
    parent = ""
    for tier, kind in enumerate(
        ["company", "function", "group", "department", "team", "subteam", "subteam"]
    ):
        ident = f"tier-{tier}"
        units[ident] = dict(
            OrgUnitID=ident,
            OrgUnitName=kind,
            ParentOrgUnitID=parent,
            UnitType=kind,
            Tier=str(tier),
            DepartmentCode="08",
        )
        parent = ident
    assert not validate(state)[0]
    db = Database("sqlite:///:memory:")
    db.save(state, 0, "test", "tiers")
    saved, _ = db.read()
    assert saved["OrganizationalUnits"]["tier-6"]["Tier"] == "6"
    assert saved["OrganizationalUnits"]["tier-3"]["DepartmentCode"] == "08"
    from backend.importer import parse_workbook

    parsed, errors, _, _ = parse_workbook(workbook_bytes(state))
    assert not errors
    assert parsed["OrganizationalUnits"]["tier-3"]["DepartmentCode"] == "08"


@pytest.mark.parametrize("tier", ["-1", "1.5", "NaN", "Infinity", "four"])
def test_invalid_explicit_tier_is_rejected(tier):
    state = sample_state()
    state["OrganizationalUnits"]["company"]["Tier"] = tier
    assert any(e["field"] == "Tier" for e in validate(state)[0])


def test_tier_type_and_parent_order_are_validated():
    state = sample_state()
    state["OrganizationalUnits"]["company"]["Tier"] = "0"
    state["OrganizationalUnits"]["eng"]["Tier"] = "1"
    assert any(e["field"] == "Tier" for e in validate(state)[0])
    state["OrganizationalUnits"]["eng"]["Tier"] = "2"
    state["OrganizationalUnits"]["platform"]["Tier"] = "3"
    state["OrganizationalUnits"]["software"]["Tier"] = "3"
    assert any("parent tier" in e["message"] for e in validate(state)[0])


def test_owner_can_create_team_only_within_assigned_scope(client):
    role(client, "owner")
    data = client.get("/api/organization").json()
    record = dict(
        OrgUnitID="owner-team",
        OrgUnitName="Team",
        ParentOrgUnitID="platform",
        UnitType="team",
        Tier="4",
    )
    saved = client.put(
        "/api/planning/OrganizationalUnits/owner-team",
        json={"version": data["version"], "record": record},
    )
    assert saved.status_code == 200, saved.text
    data = client.get("/api/organization").json()
    record.update(OrgUnitID="outside-team", ParentOrgUnitID="ops")
    denied = client.put(
        "/api/planning/OrganizationalUnits/outside-team",
        json={"version": data["version"], "record": record},
    )
    assert denied.status_code == 403
