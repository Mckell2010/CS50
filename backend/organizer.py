"""Import the organizer's unit hierarchy without guessing unassigned position mappings."""
import copy
from collections import Counter, defaultdict
from io import BytesIO
from openpyxl import load_workbook
from backend.domain import clean_value


def rows_for(ws, required):
    for n, values in enumerate(ws.iter_rows(values_only=True), 1):
        names = [clean_value(v) for v in values]
        if required <= set(names):
            return [(i, dict(zip(names, (clean_value(v) for v in values))))
                    for i, values in enumerate(ws.iter_rows(min_row=n + 1, values_only=True), n + 1)
                    if any(v is not None for v in values)]
    raise ValueError(f'{ws.title}: required column headings not found')


def normalize_entity(value):
    return {'DTI': 'divergent', 'DII': 'divergent', 'CZV': 'czinger',
            'DIVERGENT': 'divergent', 'CZINGER': 'czinger'}.get(value.strip().upper(), '')


def organizer_import(content, current, actor, deleted_units=()):
    """Return a prospective state. Caller validates and commits it atomically."""
    wb = load_workbook(BytesIO(content), read_only=True, data_only=True)
    try:
        records = rows_for(wb['Org Structure'], {'Department ID', 'Parent ID', 'Department Title', 'Tier'})
        positions = rows_for(wb['Position Control'], {'Unique Company ID', 'Department ID', 'Company'}) if 'Position Control' in wb else []
    finally:
        wb.close()
    result = copy.deepcopy(current)
    warnings, conflicts, errors = [], [], []
    units = result['OrganizationalUnits']
    incoming = {}
    for n, row in records:
        key = row.get('Department ID', '')
        if not key:
            continue
        if key in incoming:
            errors.append(dict(sheet='Org Structure', row=n, field='Department ID', message=f'Duplicate ID {key}'))
            continue
        incoming[key] = dict(OrgUnitID=key, OrgUnitName=row['Department Title'],
            ParentOrgUnitID=row.get('Parent ID', ''), DepartmentCode=key,
            DepartmentOwner=row.get('Department Owner', ''), DepartmentPurpose=row.get('Department Purpose', ''),
            Tier=row.get('Tier', ''), OwnerEmployeeID='', Entity='both', _row=n, _source='organizer')
    roots = [k for k, u in incoming.items() if not u['ParentOrgUnitID']]
    if len(roots) != 1:
        raise ValueError('Org Structure must contain exactly one Corporate root')
    root = roots[0]
    def depth(key, seen=None):
        seen = set() if seen is None else seen
        if key in seen: raise ValueError('Cycle in Org Structure parent IDs')
        seen.add(key)
        u = incoming[key]
        parent = u['ParentOrgUnitID']
        if parent and parent not in incoming: raise ValueError(f'{key}: parent {parent} is missing')
        return depth(parent, seen) + 1 if parent else 0
    names = defaultdict(list)
    for key, person in result['Employees'].items():
        names[person.get('Name', '').strip().casefold()].append(key)
    for key, u in incoming.items():
        d = depth(key)
        if u['Tier'] == '': u['Tier'] = str(d)
        tier = int(u['Tier'])
        u['UnitType'] = ['corporate', 'function', 'group', 'department', 'team'][tier] if 0 <= tier < 5 else 'subteam'
        owner_ids = names[u['DepartmentOwner'].casefold()]
        if len(owner_ids) == 1: u['OwnerEmployeeID'] = owner_ids[0]
        old = units.get(key)
        if key in deleted_units or (old and old.get('_source') == 'app'):
            conflicts.append(dict(sheet='OrganizationalUnits', id=key, message='App-owned or deleted unit preserved.'))
        else:
            units[key] = u
    if root not in units: raise ValueError('Corporate root was previously deleted; restore it in the viewer first')
    if units[root].get('ParentOrgUnitID'): raise ValueError('The Corporate root cannot have a parent')

    # Match stable person IDs only. Blank/legacy IDs do not overwrite existing assignments.
    counts = Counter(row.get('Unique Company ID') for _, row in positions)
    matched = 0
    for _, row in positions:
        key = row.get('Unique Company ID', '')
        person = result['Employees'].get(key)
        if not person or counts[key] != 1: continue
        target = row.get('Department ID', '')
        if person.get('_source') == 'app':
            if target in incoming and person.get('OrgUnitID') != target:
                conflicts.append(dict(sheet='Employees', id=key, message='App-owned assignment preserved.'))
            continue
        if target in units and target in incoming:
            person['OrgUnitID'] = target
            matched += 1
        entity = normalize_entity(row.get('Company', ''))
        if entity: person['SourceCompany'] = entity

    # Keep referenced legacy units and their ancestors, preserving budget and OKR relationships.
    used = {r['OrgUnitID'] for sheet in ['Employees', 'Budgets', 'Objectives'] for r in result[sheet].values()}
    used |= {k for k, u in units.items() if u.get('_source') == 'app'}
    for key in list(used):
        seen = set()
        while key in units and key not in seen:
            seen.add(key); used.add(key); key = units[key].get('ParentOrgUnitID')
    review = "organizer-review"
    def needs_review(key):
        seen = set()
        while key in units and key not in seen:
            if key == review: return True
            if key in incoming: return False
            seen.add(key); key = units[key].get("ParentOrgUnitID")
        return True
    legacy = {k for k in units if k != review and needs_review(k)}
    for key in legacy - used: del units[key]
    legacy &= used
    if legacy:
        units[review] = dict(OrgUnitID=review, OrgUnitName='Assignments to review', ParentOrgUnitID=root,
            UnitType='function', Tier='1', DepartmentCode='', OwnerEmployeeID='', Entity='both',
            DepartmentPurpose='Existing assignments not yet mapped to the new department IDs.', _source='organizer')
        # Retain legacy budget/OKR parent relationships within each review branch.
        for key in legacy:
            u = units[key]
            if u.get('ParentOrgUnitID') not in legacy:
                u['ParentOrgUnitID'] = review
        # Existing paths are retained as review branches; tiers reflect their new depth.
        def legacy_depth(key):
            parent = units[key].get('ParentOrgUnitID')
            return legacy_depth(parent) + 1 if parent in legacy else 2
        for key in legacy:
            d = legacy_depth(key)
            units[key]['Tier'] = str(d)
            units[key]['UnitType'] = ['corporate', 'function', 'group', 'department', 'team'][d] if d < 5 else 'subteam'
        warnings.append('Unmapped existing units and their people are retained under Assignments to review; use the table or person editor to assign the new department IDs.')
    warnings.extend([
        f'Imported {len(incoming)} corporate structure rows; {matched} existing people matched an explicit new Department ID.',
        'Blank tiers were derived from parent links. Department owner names are retained verbatim; only unique exact name matches link to a person.',
        'Position Control updates employing entity and exact new department assignments for existing stable IDs only. It does not add, deactivate, rename or re-report people; blank/old codes remain for review.',
        'All new units initially support Both entities. Use Entity in the unit editor or table to identify dedicated units; people are filtered by employing company.',
        'App-owned people and units are preserved on reimport. Budgets and objectives remain attached to their existing IDs.',
    ])
    changes=[]
    for sheet, records in result.items():
        for key, row in records.items():
            old=current[sheet].get(key)
            if old != row: changes.append(dict(sheet=sheet, id=key, kind='update' if old else 'addition', unitMove=bool(old and old.get('OrgUnitID')!=row.get('OrgUnitID'))))
        for key in set(current[sheet])-set(records):changes.append(dict(sheet=sheet,id=key,kind='removal'))
    return result, changes, conflicts, errors, warnings
