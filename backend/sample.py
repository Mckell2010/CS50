"""Explicitly synthetic fixtures and portable Excel templates."""

from io import BytesIO
from datetime import date
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from backend.domain import COLUMNS, empty_state, record_key


def sample_state(count=30):
    state = empty_state()
    year = date.today().year
    units = [
        ("company", "Divergent", "", "company", "001"),
        ("eng", "Engineering", "company", "group", "002"),
        ("ops", "Operations", "company", "group", "003"),
        ("growth", "Commercial", "company", "group", "004"),
        ("platform", "Platform & software", "eng", "department", "005"),
        ("vehicle", "Vehicle systems", "eng", "department", "006"),
        ("manufacturing", "Manufacturing", "ops", "department", "007"),
        ("people", "People & culture", "ops", "department", "008"),
        ("strategy", "Strategy & partnerships", "growth", "department", "009"),
        ("brand", "Brand & experience", "growth", "department", "010"),
        ("software", "Digital engineering", "platform", "subdepartment", "011"),
    ]
    for ident, name, parent, kind, owner in units:
        state["OrganizationalUnits"][ident] = dict(
            zip(COLUMNS["OrganizationalUnits"], [ident, name, parent, kind, owner]),
            _source="excel",
        )
    names = [
        "Lukas Czinger",
        "Alex Morgan",
        "Jordan Lee",
        "Taylor Brooks",
        "Maya Chen",
        "Noah Williams",
        "Avery Patel",
        "Sofia Rivera",
        "Elliot Hayes",
        "Olivia Park",
        "Riley Bennett",
        "Amara Okafor",
        "Theo Martin",
        "Isla Thompson",
        "Kai Robinson",
        "Zara Wilson",
        "Leo Garcia",
        "Nina Ahmed",
        "Sam Reed",
        "Ari Collins",
        "Eva Kim",
        "Jules Walker",
        "Luca Rossi",
        "Mila Cooper",
        "Rowan Blake",
        "Ivy Santos",
        "Owen Price",
        "Remy Scott",
        "Ada Lewis",
        "Finn Young",
    ]
    titles = [
        "Chief Executive Officer",
        "VP, Engineering",
        "VP, Operations",
        "VP, Commercial",
        "Director, Platform",
        "Director, Vehicle Systems",
        "Director, Manufacturing",
        "Director, People",
        "Director, Strategy",
        "Director, Brand",
        "Engineering Manager",
    ]
    for i in range(count):
        ident = f"{i+1:03d}"
        if i == 0:
            manager = ""
            unit = "company"
        elif i < 4:
            manager = "001"
            unit = ["eng", "ops", "growth"][i - 1]
        elif i < 10:
            manager = ["002", "002", "003", "003", "004", "004"][i - 4]
            unit = units[i][0]
        elif i == 10:
            manager = "005"
            unit = "software"
        else:
            branch = (i - 11) % 6
            manager = ["011", "006", "007", "008", "009", "010"][branch]
            unit = [
                "software",
                "vehicle",
                "manufacturing",
                "people",
                "strategy",
                "brand",
            ][branch]
        # Correct department mapping uses index 4..9 in units.
        row = {
            "EmployeeID": ident,
            "Name": names[i] if i < len(names) else f"Sample Person {i+1}",
            "JobTitle": (
                titles[i]
                if i < len(titles)
                else [
                    "Software Engineer",
                    "Systems Engineer",
                    "Manufacturing Specialist",
                    "People Partner",
                    "Strategy Associate",
                    "Brand Designer",
                ][(i - 11) % 6]
            ),
            "ManagerID": manager,
            "OrgUnitID": unit,
            "EmploymentStatus": "active",
            "FTE": "0.8" if i == 20 else "1",
            "StartDate": "2025-01-01",
            "EndDate": "",
            "Location": "Torrance, CA" if i % 3 else "Remote",
            "EmploymentType": "Full time",
            "_source": "excel",
        }
        state["Employees"][ident] = row
        c = {
            "EmployeeID": ident,
            "EffectiveFrom": "2025-01-01",
            "EffectiveTo": "",
            "Currency": "USD",
            "AnnualBasePay": str(
                300000
                if i == 0
                else 220000 if i < 4 else 165000 if i < 11 else 105000 + (i % 5) * 7000
            ),
            "AnnualTargetBonus": "20000" if i < 11 else "10000",
            "AnnualEmployerCosts": "25000",
            "_source": "excel",
        }
        state["Compensation"][record_key("Compensation", c)] = c
    envelopes = {
        "company": 7500000,
        "eng": 2700000,
        "ops": 2100000,
        "growth": 2100000,
        "platform": 1200000,
        "vehicle": 1100000,
        "manufacturing": 850000,
        "people": 850000,
        "strategy": 850000,
        "brand": 850000,
        "software": 850000,
    }
    for ident, _, _, _, _ in units:
        r = {
            "BudgetID": f"budget-{ident}-{year}",
            "OrgUnitID": ident,
            "BudgetYear": str(year),
            "Currency": "USD",
            "ApprovedEnvelope": str(envelopes[ident]),
            "DirectAllocation": str(
                300000
                if ident in ["company", "eng", "ops", "growth", "platform"]
                else envelopes[ident]
            ),
            "_source": "excel",
        }
        state["Budgets"][r["BudgetID"]] = r
    objectives = [
        (
            "o-company",
            "company",
            "001",
            "",
            "Build the next chapter of sustainable performance",
            "on-track",
        ),
        (
            "o-eng",
            "eng",
            "002",
            "o-company",
            "Bring the next platform to production readiness",
            "on-track",
        ),
        (
            "o-ops",
            "ops",
            "003",
            "o-company",
            "Create a repeatable manufacturing system",
            "at-risk",
        ),
        (
            "o-growth",
            "growth",
            "004",
            "o-company",
            "Deepen our strategic partnerships",
            "on-track",
        ),
    ]
    for ident, unit, owner, parent, text, status in objectives:
        state["Objectives"][ident] = {
            "ObjectiveID": ident,
            "OrgUnitID": unit,
            "OwnerEmployeeID": owner,
            "ParentObjectiveID": parent,
            "ObjectiveText": text,
            "PeriodStart": f"{year}-01-01",
            "PeriodEnd": f"{year}-12-31",
            "Status": status,
            "_source": "excel",
        }
    krs = [
        (
            "kr1",
            "o-eng",
            "002",
            "Validation milestones completed",
            "milestones",
            "0",
            "7",
            "10",
            "increase",
        ),
        (
            "kr2",
            "o-ops",
            "003",
            "Manufacturing cycle time",
            "hours",
            "40",
            "32",
            "25",
            "decrease",
        ),
        (
            "kr3",
            "o-growth",
            "004",
            "Active strategic partnerships",
            "partners",
            "2",
            "6",
            "8",
            "increase",
        ),
    ]
    for ident, obj, owner, metric, unit, b, c, t, d in krs:
        state["KeyResults"][ident] = {
            "KeyResultID": ident,
            "ObjectiveID": obj,
            "OwnerEmployeeID": owner,
            "MetricName": metric,
            "Unit": unit,
            "Baseline": b,
            "CurrentValue": c,
            "TargetValue": t,
            "Direction": d,
            "LastUpdated": date.today().isoformat(),
            "_source": "excel",
        }
    return state


def workbook_bytes(state=None):
    wb = Workbook()
    info = wb.active
    info.title = "Read me"
    notes = [
        ["DIVERGENT • ORGANIZATION WORKBOOK"],
        ["SYNTHETIC SAMPLE DATA" if state else "EMPTY IMPORT TEMPLATE"],
        [
            "Only the example CEO name is provided by the brief; all other people, pay, relationships, and objectives are invented."
        ],
        [
            "IDs are text. Preserve leading zeros. Dates use YYYY-MM-DD. Currency defaults to USD; annual budgets use calendar years."
        ],
        [
            "Pay fields are actual annual amounts at contracted FTE. Do not multiply by FTE again. Blank pay means unknown."
        ],
        ["Employees and OrganizationalUnits are required. Other sheets are optional."],
        ["EffectiveTo and EndDate are inclusive. Compensation periods cannot overlap."],
        [
            "Import defaults to incremental updates; omitted rows remain. Full snapshot explicitly deactivates missing employees."
        ],
        [
            "App-owned planning records are preserved. Resolve conflicts in the planning editor."
        ],
        [
            "EmploymentStatus: active / leave / inactive. UnitType: company / group / department / subdepartment."
        ],
        [
            "Objective Status: on-track / at-risk / off-track / complete / not-started. Direction: increase / decrease."
        ],
        [
            "ApprovedEnvelope covers a unit and all descendants. DirectAllocation reserves budget for directly assigned people."
        ],
        ["No objective score is calculated by averaging unrelated key results."],
        [
            "All IDs, references, dates, FTE, pay, owners, hierarchies, and intervals are validated before commit."
        ],
    ]
    for row in notes:
        info.append(row)
    info.column_dimensions["A"].width = 125
    for sheet, columns in COLUMNS.items():
        ws = wb.create_sheet(sheet)
        ws.append(columns)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{ws.cell(1,len(columns)).column_letter}1"
        if state:
            for record in state[sheet].values():
                ws.append([record.get(c, "") for c in columns])
        for cell in ws[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = PatternFill("solid", fgColor="25483D")
        for idx, col in enumerate(columns, 1):
            ws.column_dimensions[ws.cell(1, idx).column_letter].width = max(
                20, len(col) + 3
            )
            if col.endswith("ID"):
                for cells in ws.iter_rows(min_row=2, min_col=idx, max_col=idx):
                    cells[0].number_format = "@"
    out = BytesIO()
    wb.save(out)
    return out.getvalue()
