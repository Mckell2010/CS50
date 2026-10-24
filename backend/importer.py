import json, hashlib, zipfile
from io import BytesIO
from openpyxl import load_workbook
from backend.domain import COLUMNS, REQUIRED, clean_value, record_key, empty_state

MAX_UPLOAD = 10 * 1024 * 1024


def parse_workbook(content, mapping=None):
    if len(content) > MAX_UPLOAD:
        raise ValueError("Workbook exceeds the 10 MB upload limit")
    try:
        with zipfile.ZipFile(BytesIO(content)) as z:
            if sum(i.file_size for i in z.infolist()) > 60 * 1024 * 1024:
                raise ValueError("Expanded workbook exceeds 60 MB")
        wb = load_workbook(BytesIO(content), read_only=True, data_only=False)
    except (zipfile.BadZipFile, KeyError, ValueError) as e:
        raise ValueError(f"Invalid .xlsx workbook: {e}")
    if "Org Structure" in wb.sheetnames:
        wb.close()
        from backend.organizer import organizer_import
        state, _, _, errors, warnings = organizer_import(content, empty_state(), "import")
        return state, errors, {"Org Structure": []}, warnings
    if "Employees" not in wb.sheetnames:
        from backend.chart_import import chart_export, parse_chart

        detected = chart_export(wb)
        if detected:
            result = parse_chart(*detected)
            wb.close()
            return result
    mapping = mapping or {}
    state = empty_state()
    errors = []
    headers = {}
    warnings = []
    for sheet in ["Employees", "OrganizationalUnits"]:
        if sheet not in wb.sheetnames:
            errors.append(
                {
                    "sheet": sheet,
                    "row": 1,
                    "field": "sheet",
                    "message": "Required sheet is missing",
                }
            )
    for sheet, columns in COLUMNS.items():
        if sheet not in wb.sheetnames:
            continue
        ws = wb[sheet]
        rows = ws.iter_rows(values_only=True)
        names = [clean_value(c) for c in next(rows, [])]
        headers[sheet] = names
        selected = {
            field: mapping.get(sheet, {}).get(field, field) for field in columns
        }
        indexes = {
            field: names.index(name)
            for field, name in selected.items()
            if name in names
        }
        for field in REQUIRED[sheet]:
            if sheet == "Budgets" and (
                (field == "ApprovedEnvelope" and "AllocationPercent" in indexes)
                or (field == "DirectAllocation" and "DirectPercent" in indexes)
            ):
                continue
            if (
                sheet == "Compensation"
                and "PayBasis" in indexes
                and field
                in {"AnnualBasePay", "AnnualTargetBonus", "AnnualEmployerCosts"}
            ):
                continue
            if field not in indexes:
                errors.append(
                    {
                        "sheet": sheet,
                        "row": 1,
                        "field": field,
                        "message": "Column missing. Map an existing column to this field.",
                    }
                )
        if len([x for x in names if x]) != len(set(x for x in names if x)):
            errors.append(
                {
                    "sheet": sheet,
                    "row": 1,
                    "field": "headers",
                    "message": "Duplicate column headers",
                }
            )
        for rownum, values in enumerate(rows, 2):
            if rownum > 20001:
                raise ValueError("Maximum 20,000 rows per sheet")
            if not any(v is not None for v in values):
                continue
            record = {
                field: clean_value(values[index]) if index < len(values) else ""
                for field, index in indexes.items()
            }
            for field in columns:
                record.setdefault(field, "")
            record["_row"] = rownum
            for field, value in record.items():
                if isinstance(value, str) and value.startswith("="):
                    errors.append(
                        {
                            "sheet": sheet,
                            "row": rownum,
                            "field": field,
                            "message": "Use literal values instead of Excel formulas",
                        }
                    )
            key = record_key(sheet, record)
            if key in state[sheet]:
                errors.append(
                    {
                        "sheet": sheet,
                        "row": rownum,
                        "field": columns[0],
                        "message": f"Duplicate ID {key}",
                    }
                )
            state[sheet][key] = record
    wb.close()
    return state, errors, headers, warnings
