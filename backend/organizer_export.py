"""Organizer-layout export plus stable-ID tables for Power BI. No inferred pay or approvals."""
from collections import Counter
from datetime import date, datetime, timezone
from io import BytesIO
import re
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from backend.domain import COLUMNS, clean_value
from backend.organizer import normalize_entity, rows_for

ORG_HEADERS = ['Department ID', 'Parent ID', 'Department Title', 'Department Owner', 'Department Purpose', 'Tier']
POSITION_HEADERS = [
    'Unique Company ID', 'Trinet ID', 'Position ID', 'Position ID Status', 'Role Type',
    'Company', 'Location', 'Department ID', 'Position', 'Date Hired / Date Needed',
    'Employment Type', 'EEO Job Category', 'Name', 'Direct Manager', 'Occupation', 'Level',
    'Temp in Role', 'Priority', 'Stage', 'Recruiter', 'Cycle Time \n(PID Approved to PID Fill)',
    'Filled\nDate', 'Need Date', 'PID Approval Date', 'Inactivated Date', 'Hire Date',
    'Recruitment On-Time Flag', 'Req ID', 'Finance Budgeted Position Title',
    'Hire Date / Need Date Month', 'Tier', 'Desk Required',
]
DATE_COLUMNS = {'Date Hired / Date Needed', 'Filled\nDate', 'Need Date', 'PID Approval Date',
                'Inactivated Date', 'Hire Date', 'StartDate', 'EndDate'}
ILLEGAL_XML = re.compile(r'[\x00-\x08\x0b\x0c\x0e-\x1f]')


def organizer_source(content):
    """Retain literal/cached source fields, not workbook formulas or executable content."""
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    try:
        if 'Position Control' not in workbook:
            return {'headers': POSITION_HEADERS, 'positions': []}
        sheet = workbook['Position Control']
        rows = rows_for(sheet, {'Unique Company ID', 'Department ID', 'Company'})
        headers = next(([clean_value(v) for v in values if clean_value(v)]
                        for values in sheet.iter_rows(values_only=True)
                        if {'Unique Company ID', 'Department ID', 'Company'} <= {clean_value(v) for v in values}), POSITION_HEADERS)
        if len(set(h.casefold() for h in headers)) != len(headers):
            raise ValueError('Position Control contains duplicate column headings')
        if len(rows) > 20000:
            raise ValueError('Maximum 20,000 position rows')
        omitted = [h for h in headers if h not in POSITION_HEADERS]
        headers = [h for h in headers if h in POSITION_HEADERS]
        return {'headers': headers, 'omittedColumnCount': len(omitted), 'positions': [dict(sourceRow=n, values={h: row.get(h, '') for h in headers}) for n, row in rows]}
    finally:
        workbook.close()


def literal(cell, value):
    if isinstance(value, str):
        cell.value = ILLEGAL_XML.sub('', value)
        cell.data_type = 's'  # Prevent formula injection, including user-entered names/purposes.
    else:
        cell.value = value


def write_table(workbook, title, headers, rows, table_name, header_row=1):
    sheet = workbook.create_sheet(title)
    if header_row > 1:
        literal(sheet.cell(1, 1), title)
        sheet.cell(1, 1).font = Font(name='Arial', size=18, bold=True)
    for col, header in enumerate(headers, 1):
        cell = sheet.cell(header_row, col)
        literal(cell, header)
        cell.fill = PatternFill('solid', fgColor='151515')
        cell.font = Font(name='Arial', color='FFFFFF', bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical='center')
    sheet.row_dimensions[header_row].height = 36
    for row_index, row in enumerate(rows, header_row + 1):
        for col, (header, value) in enumerate(zip(headers, row), 1):
            cell = sheet.cell(row_index, col)
            if title in {'Org Structure', 'Org Units', 'People'} and header in {'Tier', 'FTE'} and value not in ('', None):
                value = int(float(value)) if header == 'Tier' else float(value)
            if header in DATE_COLUMNS and isinstance(value, str) and re.fullmatch(r'\d{4}-\d{2}-\d{2}', value):
                try:
                    value = date.fromisoformat(value)
                    cell.number_format = 'yyyy-mm-dd'
                except ValueError:
                    pass  # Preserve invalid source dates as text for review, rather than inventing a date.
            literal(cell, value if value != '' else None)
            if isinstance(value, str):
                cell.number_format = '@'
            cell.alignment = Alignment(vertical='top', wrap_text=True)
    if rows:
        from openpyxl.utils import get_column_letter
        ref = f'A{header_row}:{get_column_letter(len(headers))}{header_row + len(rows)}'
        table = Table(displayName=table_name, ref=ref)
        table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
        sheet.add_table(table)
    for col, header in enumerate(headers, 1):
        from openpyxl.utils import get_column_letter
        width = 38 if header in {'Department Purpose', 'OrgUnitName', 'Department Title', 'Position', 'Review note'} else 23
        sheet.column_dimensions[get_column_letter(col)].width = width
    sheet.freeze_panes = f'A{header_row + 1}'
    return sheet


def organizer_workbook(state, source=None, version=0, synthetic=False):
    source = source or {}
    people = state['Employees']
    units = state['OrganizationalUnits']
    notes = []
    def review(kind, key, message):
        notes.append([kind, key, message])
    if not source:
        review('Source', '', 'Original Position Control fields have not been retained yet. Reimport the original organizer to preserve recruiting and approval fields; only known app data is exported.')
    if synthetic:
        review('Dataset', '', 'SYNTHETIC SAMPLE DATA — not the actual organization.')
    review('Snapshot', str(version), f'Exported {datetime.now(timezone.utc).isoformat()} (UTC). Includes both entities, all current app records and source-only position rows.')
    review('Power BI', '', 'Use tblOrgUnits and tblPeople as the current app snapshot. Join People.OrgUnitID → OrgUnits.OrgUnitID. ManagerID is a separate relationship to People.EmployeeID. Treat IDs as text.')
    review('Position Control', '', 'Original columns are preserved. Current app fields update exact stable-ID matches; other source rows remain unchanged and are marked Source only in Position Links. Do not union this sheet with People or count them as two rosters.')
    review('Tier fields', '', 'Org Structure / Org Units Tier is hierarchy depth. Position Control Tier is retained from the source workbook and is not replaced with hierarchy depth.')
    review('Completeness', '', 'Unknown values stay blank. Vacancy does not imply approval. Compensation and performance data are not included. Formatting is recreated; formulas/macros are not carried over.')
    if source.get('filename'):
        review('Source', source['filename'], f"Last organizer import version: {source.get('importVersion', 'unknown')}")
    org_rows=[]
    def order_key(unit):
        path=[]; key=unit['OrgUnitID'];seen=set()
        while key in units and key not in seen:
            seen.add(key);path.insert(0,key);key=units[key].get('ParentOrgUnitID')
        return tuple(path)
    for u in sorted(units.values(), key=order_key):
        owner=people.get(u.get('OwnerEmployeeID'),{}).get('Name','') or u.get('DepartmentOwner','')
        tier=u.get('Tier','')
        org_rows.append([u['OrgUnitID'],u.get('ParentOrgUnitID',''),u['OrgUnitName'],owner,u.get('DepartmentPurpose',''),int(float(tier)) if tier != '' else None])
        if not owner: review('Unit',u['OrgUnitID'],'Department owner is not assigned.')
        if not u.get('DepartmentPurpose'): review('Unit',u['OrgUnitID'],'Department purpose is not supplied.')
    if source.get('omittedColumnCount'):
        review('Source', '', f"{source['omittedColumnCount']} nonstandard source columns were excluded. Only the original organizer fields are carried forward.")
    headers=[h for h in (source.get('headers') or POSITION_HEADERS) if h in POSITION_HEADERS]
    # Required original-layout columns remain available even in a minimal source workbook.
    headers=list(headers)+[h for h in POSITION_HEADERS if h not in headers]
    source_rows=source.get('positions',[])
    ids=Counter(r['values'].get('Unique Company ID','') for r in source_rows)
    pids=Counter(r['values'].get('Position ID','') for r in source_rows)
    exported=[]; links=[]; matched=set(); position_ids={}
    def update(row, person):
        row=dict(row)
        unit=units[person['OrgUnitID']]
        employment=person.get('EmploymentStatus','')
        row.update({'Department ID':person['OrgUnitID'], 'Position':person.get('JobTitle',''),
                    'Name':'' if employment=='vacant' else person.get('Name',''),
                    'Direct Manager':people.get(person.get('ManagerID'),{}).get('Name',''),
                    'Employment Type':'Full Time' if person.get('EmploymentType')=='FTE' else person.get('EmploymentType','')})
        company=normalize_entity(person.get('SourceCompany',''))
        if company:
            row['Company'] = row.get('Company') if normalize_entity(row.get('Company',''))==company else ('DTI' if company=='divergent' else 'CZV')
        if person.get('Location'):
            row['Location']={'Los Angeles - HQ':'HQ','Los Angeles - LB':'LB'}.get(person['Location'],person['Location'])
        if employment in {'active','leave'}: row['Position ID Status']='Active'
        elif employment=='inactive': row['Position ID Status']='Inactive'
        elif row.get('Position ID Status') in {'Active','Filled','Inactive'}:
            row['Position ID Status']=''
            review('Person',person['EmployeeID'],'Source position status conflicts with the current vacancy/draft. Approval was not inferred; check Position ID Status.')
        if person.get('StartDate') and employment not in {'vacant','draft'}:
            row['Date Hired / Date Needed']=person['StartDate'];row['Hire Date']=person['StartDate']
        if person.get('EndDate') and employment=='inactive': row['Inactivated Date']=person['EndDate']
        if unit['OrgUnitID']=='organizer-review' or 'organizer-review' in order_key(unit):
            review('Person',person['EmployeeID'],'Department assignment is still in Assignments to review.')
        return row
    for item in source_rows:
        row=dict(item['values']); ident=row.get('Unique Company ID','');pid=row.get('Position ID','')
        key=ident if ids[ident]==1 and ident in people else ''
        # Existing chart-export vacancy IDs are the literal Position ID plus -Open.
        vacancy=f'{pid}-Open'
        if not key and pid and pids[pid]==1 and vacancy in people and people[vacancy].get('EmploymentStatus')=='vacant': key=vacancy
        if key in matched: key=''
        if key:
            row=update(row,people[key]);matched.add(key);position_ids[key]=pid
        else:
            review('Source row',str(item.get('sourceRow','')),'No unique current app match; original Position Control values retained.')
        exported.append([row.get(h,'') for h in headers])
        links.append([len(exported)+3,item.get('sourceRow',''),key,pid,'Matched app record' if key else 'Source only'])
    for key,person in people.items():
        if key in matched: continue
        row=update({'Unique Company ID':key},person)
        exported.append([row.get(h,'') for h in headers])
        links.append([len(exported)+3,'',key,'','App only'])
        review('Person',key,'No uniquely matched source position. App record appended; recruiting/approval fields remain blank. Review possible duplicates against Source only rows.')
    wb=Workbook();wb.remove(wb.active)
    write_table(wb,'Org Structure',ORG_HEADERS,org_rows,'tblOrgStructure',4)
    write_table(wb,'Position Control',headers,exported,'tblPositionControl',3)
    unit_headers=COLUMNS['OrganizationalUnits']
    write_table(wb,'Org Units',unit_headers,[[u.get(h,'') for h in unit_headers] for u in sorted(units.values(),key=order_key)],'tblOrgUnits')
    people_headers=COLUMNS['Employees']+['PositionID']
    write_table(wb,'People',people_headers,[[p.get(h,'') for h in COLUMNS['Employees']]+[position_ids.get(k,'')] for k,p in people.items()],'tblPeople')
    write_table(wb,'Position Links',['Export Excel Row','Original Excel Row','EmployeeID','PositionID','Match status'],links,'tblPositionLinks')
    write_table(wb,'Export Notes',['Record type','Record ID','Review note'],notes,'tblExportNotes')
    wb.properties.title='Org Chart Organizer — current app snapshot'
    result=BytesIO();wb.save(result);return result.getvalue()
